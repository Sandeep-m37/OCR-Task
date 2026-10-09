import os
import re
import tempfile
from datetime import datetime

import pymupdf
import pytesseract

from PIL import Image
from pytesseract import Output

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


# =========================================================
# SETTINGS
# =========================================================

INPUT_FILE = "bill1.jpg"
OUTPUT_FILE = "generic_bill.xlsx"


# =========================================================
# OCR - IMAGE
# =========================================================

def extract_text_from_image(image_path):

    image = Image.open(image_path)

    text = pytesseract.image_to_string(
        image,
        config="--psm 6"
    )

    return text


# =========================================================
# OCR WORD POSITION - IMAGE
# =========================================================

def extract_words_from_image(image_path):

    image = Image.open(image_path)

    data = pytesseract.image_to_data(
        image,
        config="--psm 6",
        output_type=Output.DICT
    )

    words = []

    for i in range(len(data["text"])):

        word = data["text"][i].strip()

        if not word:
            continue

        try:
            confidence = float(data["conf"][i])
        except:
            confidence = 0

        if confidence < 20:
            continue

        words.append({
            "text": word,
            "left": data["left"][i],
            "top": data["top"][i],
            "width": data["width"][i],
            "height": data["height"][i]
        })

    return words


# =========================================================
# PDF -> IMAGES
# =========================================================

def pdf_to_images(pdf_path):

    pdf = pymupdf.open(pdf_path)

    temp_folder = tempfile.mkdtemp()

    image_paths = []

    for page_number, page in enumerate(pdf):

        pix = page.get_pixmap(
            matrix=pymupdf.Matrix(3, 3),
            alpha=False
        )

        image_path = os.path.join(
            temp_folder,
            f"page_{page_number + 1}.png"
        )

        pix.save(image_path)

        image_paths.append(image_path)

    pdf.close()

    return image_paths


# =========================================================
# PDF DIRECT TEXT + WORD POSITION
# =========================================================

def extract_text_and_words_from_pdf(pdf_path):

    pdf = pymupdf.open(pdf_path)

    all_text = []
    all_words = []

    for page in pdf:

        # -------------------------------------------------
        # DIRECT PDF TEXT
        # -------------------------------------------------

        text = page.get_text("text")

        if text.strip():
            all_text.append(text)

        # -------------------------------------------------
        # DIRECT PDF WORD POSITION
        # -------------------------------------------------

        page_words = page.get_text("words")

        for item in page_words:

            x0 = item[0]
            y0 = item[1]
            x1 = item[2]
            y1 = item[3]
            word = item[4].strip()

            if not word:
                continue

            all_words.append({
                "text": word,
                "left": x0,
                "top": y0,
                "width": x1 - x0,
                "height": y1 - y0
            })

    pdf.close()

    return "\n".join(all_text), all_words


# =========================================================
# MAIN FILE READER
# =========================================================

def extract_bill_data(file_path):

    extension = os.path.splitext(file_path)[1].lower()

    image_extensions = [
        ".jpg",
        ".jpeg",
        ".png",
        ".bmp",
        ".tiff",
        ".tif",
        ".webp"
    ]

    # =====================================================
    # IMAGE
    # =====================================================

    if extension in image_extensions:

        text = extract_text_from_image(file_path)

        words = extract_words_from_image(file_path)

        return text, words

    # =====================================================
    # PDF
    # =====================================================

    elif extension == ".pdf":

        # -------------------------------------------------
        # FIRST TRY: DIRECT PDF TEXT
        # -------------------------------------------------

        text, words = extract_text_and_words_from_pdf(
            file_path
        )

        # -------------------------------------------------
        # IF PDF HAS GOOD TEXT
        # -------------------------------------------------

        if len(text.strip()) >= 30:

            print(
                "\nPDF type: Digital/Text PDF"
            )

            print(
                "Using direct PDF text extraction..."
            )

            return text, words

        # -------------------------------------------------
        # OTHERWISE OCR
        # -------------------------------------------------

        print(
            "\nPDF type: Scanned/Image PDF"
        )

        print(
            "Using Tesseract OCR..."
        )

        image_paths = pdf_to_images(file_path)

        all_text = []
        all_words = []

        for image_path in image_paths:

            page_text = extract_text_from_image(
                image_path
            )

            page_words = extract_words_from_image(
                image_path
            )

            all_text.append(page_text)

            all_words.extend(page_words)

        return "\n".join(all_text), all_words

    else:

        raise ValueError(
            "Use JPG, JPEG, PNG, BMP, TIFF, WEBP or PDF."
        )


# =========================================================
# TEXT CLEANING
# =========================================================

def clean_text(text):

    lines = []

    for line in text.splitlines():

        line = line.strip()

        if not line:
            continue

        line = re.sub(
            r"\s+",
            " ",
            line
        )

        lines.append(line)

    return lines


# =========================================================
# DATE
# =========================================================

DATE_PATTERN = (
    r"\b(?:"
    r"\d{1,2}[/-]\d{1,2}[/-]\d{2,4}"
    r"|"
    r"\d{4}[/-]\d{1,2}[/-]\d{1,2}"
    r"|"
    r"\d{1,2}\s+[A-Za-z]{3,9}\s+\d{2,4}"
    r"|"
    r"[A-Za-z]{3,9}\s+\d{1,2},\s+\d{2,4}"
    r")\b"
)


def find_date_in_line(line):

    match = re.search(
        DATE_PATTERN,
        line,
        re.IGNORECASE
    )

    if match:
        return match.group(0)

    return None


# =========================================================
# PHONE
# =========================================================

PHONE_PATTERN = re.compile(
    r"(?<!\d)"
    r"(?:\+\d{1,3}[\s.-]?)?"
    r"(?:\(\d{2,5}\)[\s.-]?)?"
    r"\d{3,5}[\s.-]\d{3,5}"
    r"(?!\d)"
)


def find_phone_in_line(line):

    matches = PHONE_PATTERN.findall(line)

    for phone in matches:

        digits = re.sub(
            r"\D",
            "",
            phone
        )

        if 7 <= len(digits) <= 15:
            return phone.strip()

    return None


def find_phone(lines):

    for line in lines:

        phone = find_phone_in_line(line)

        if phone:
            return phone

    return None


# =========================================================
# INVOICE NUMBER
# =========================================================

def find_invoice_number(lines):

    patterns = [

        r"(?:invoice\s*(?:no|number|#))\s*[:#-]?\s*"
        r"([A-Za-z0-9][A-Za-z0-9./_-]*)",

        r"(?:inv\s*(?:no|number|#))\s*[:#-]?\s*"
        r"([A-Za-z0-9][A-Za-z0-9./_-]*)",

        r"(?:bill\s*(?:no|number|#))\s*[:#-]?\s*"
        r"([A-Za-z0-9][A-Za-z0-9./_-]*)"
    ]

    for line in lines:

        for pattern in patterns:

            match = re.search(
                pattern,
                line,
                re.IGNORECASE
            )

            if match:

                value = match.group(1).strip()

                if value:
                    return value

    return None


# =========================================================
# INVOICE DATE
# =========================================================

def find_invoice_date(lines):

    labels = [
        "invoice date",
        "bill date",
        "date of invoice",
        "invoice dt"
    ]

    for line in lines:

        lower_line = line.lower()

        for label in labels:

            if label in lower_line:

                date = find_date_in_line(line)

                if date:
                    return date

    return None


# =========================================================
# AMOUNT
# =========================================================

AMOUNT_PATTERN = (
    r"(?:₹|Rs\.?|INR|USD|\$|EUR|€|GBP|£)?\s*"
    r"\d[\d,]*(?:\.\d{1,2})?"
)


def find_amounts(text):

    matches = re.findall(
        AMOUNT_PATTERN,
        text,
        re.IGNORECASE
    )

    amounts = []

    for value in matches:

        number = re.search(
            r"\d[\d,]*(?:\.\d{1,2})?",
            value
        )

        if number:

            amounts.append(
                number.group()
            )

    return amounts


# =========================================================
# SUBTOTAL
# =========================================================

def find_subtotal(lines):

    labels = [
        "subtotal",
        "sub total",
        "sub-total"
    ]

    for line in lines:

        lower_line = line.lower()

        for label in labels:

            if label in lower_line:

                amounts = find_amounts(line)

                if amounts:
                    return amounts[-1]

    return None


# =========================================================
# TAX
# =========================================================

def find_tax(lines):

    values = []

    labels = [
        "total tax",
        "gst",
        "vat",
        "cgst",
        "sgst",
        "igst",
        "service tax",
        "tax"
    ]

    for line in lines:

        lower_line = line.lower()

        if "tax invoice" in lower_line:
            continue

        for label in labels:

            if label in lower_line:

                amounts = find_amounts(line)

                if amounts:

                    values.append(
                        amounts[-1]
                    )

                break

    if values:
        return ", ".join(values)

    return None


# =========================================================
# DISCOUNT
# =========================================================

def find_discount(lines):

    labels = [
        "discount",
        "coupon",
        "promo",
        "offer"
    ]

    for line in lines:

        lower_line = line.lower()

        for label in labels:

            if label in lower_line:

                amounts = find_amounts(line)

                if amounts:
                    return amounts[-1]

    return None


# =========================================================
# OTHER CHARGES
# =========================================================

def find_other_charges(lines):

    labels = [
        "delivery charge",
        "shipping charge",
        "platform fee",
        "service charge",
        "convenience fee",
        "handling charge",
        "booking fee",
        "processing fee"
    ]

    values = []

    for line in lines:

        lower_line = line.lower()

        for label in labels:

            if label in lower_line:

                amounts = find_amounts(line)

                if amounts:

                    values.append(
                        amounts[-1]
                    )

                break

    if values:
        return ", ".join(values)

    return None


# =========================================================
# FINAL TOTAL
# =========================================================

def find_total(lines):

    strong_labels = [
        "grand total",
        "amount payable",
        "amount due",
        "net payable",
        "final total",
        "total payable",
        "invoice total",
        "balance due"
    ]

    # -----------------------------------------------------
    # FIRST - STRONG TOTAL LABEL
    # -----------------------------------------------------

    for line in reversed(lines):

        lower_line = line.lower()

        if "subtotal" in lower_line:
            continue

        for label in strong_labels:

            if label in lower_line:

                amounts = find_amounts(line)

                if amounts:
                    return amounts[-1]

    # -----------------------------------------------------
    # SECOND - NORMAL TOTAL
    # -----------------------------------------------------

    for line in reversed(lines):

        lower_line = line.lower()

        if "subtotal" in lower_line:
            continue

        if "total tax" in lower_line:
            continue

        if re.search(
            r"\btotal\b",
            lower_line
        ):

            amounts = find_amounts(line)

            if amounts:
                return amounts[-1]

    return None


# =========================================================
# BILL INFORMATION LABELS
# =========================================================

FIELD_LABELS = {

    "GSTIN": [
        "gstin",
        "gst no",
        "gst number"
    ],

    "Customer Name": [
        "customer name",
        "buyer name",
        "guest name",
        "patient name",
        "passenger name"
    ],

    "Address": [
        "billing address",
        "shipping address",
        "customer address"
    ],

    "Payment Method": [
        "payment method",
        "payment mode",
        "mode of payment",
        "paid by"
    ],

    "Due Date": [
        "due date",
        "payment due"
    ],

    "Check In": [
        "check in",
        "check-in",
        "checkin"
    ],

    "Check Out": [
        "check out",
        "check-out",
        "checkout"
    ],

    "Start Date": [
        "start date",
        "service start",
        "from date"
    ],

    "End Date": [
        "end date",
        "service end",
        "to date"
    ],

    "Billing Period": [
        "billing period",
        "service period",
        "subscription period"
    ],

    "Order Date": [
        "order date"
    ],

    "Order ID": [
        "order id",
        "order number",
        "order no"
    ],

    "Transaction ID": [
        "transaction id",
        "transaction number",
        "txn id",
        "reference id"
    ],

    "From": [
        "pickup location",
        "pickup",
        "pick up"
    ],

    "To": [
        "drop location",
        "drop off",
        "drop"
    ],

    "Distance": [
        "distance"
    ],

    "Fare": [
        "fare",
        "ride fare",
        "trip fare"
    ],

    "Duration": [
        "duration",
        "trip duration"
    ],

    "Nights": [
        "nights",
        "number of nights"
    ]
}


# =========================================================
# EXTRACT KNOWN FIELDS
# =========================================================

def extract_known_fields(lines):

    fields = {}

    for field_name, aliases in FIELD_LABELS.items():

        for line in lines:

            lower_line = line.lower()

            for alias in aliases:

                if alias in lower_line:

                    pattern = (
                        re.escape(alias)
                        + r"\s*[:#-]\s*(.+)"
                    )

                    match = re.search(
                        pattern,
                        line,
                        re.IGNORECASE
                    )

                    if match:

                        value = match.group(1).strip()

                        if value:

                            fields[field_name] = value

                            break

            if field_name in fields:
                break

    return fields


# =========================================================
# DYNAMIC LABELS
# =========================================================

def find_dynamic_labels(lines):

    fields = {}

    ignored = [
        "item",
        "description",
        "amount",
        "total",
        "tax",
        "qty",
        "quantity",
        "price",
        "rate"
    ]

    for line in lines:

        match = re.match(
            r"^\s*([A-Za-z][A-Za-z0-9 /_-]{1,40})"
            r"\s*:\s*(.+)$",
            line
        )

        if not match:
            continue

        label = match.group(1).strip()

        value = match.group(2).strip()

        if any(
            word in label.lower()
            for word in ignored
        ):
            continue

        if value:

            fields[label.title()] = value

    return fields


# =========================================================
# VENDOR
# =========================================================

def find_vendor(lines):

    ignored = [
        "invoice",
        "tax invoice",
        "bill",
        "receipt",
        "order summary",
        "statement"
    ]

    for line in lines[:10]:

        lower_line = line.lower()

        if any(
            word in lower_line
            for word in ignored
        ):
            continue

        if find_phone_in_line(line):
            continue

        if find_date_in_line(line):
            continue

        if len(line) >= 3:
            return line

    return None


# =========================================================
# BILL FIELDS
# =========================================================

def create_bill_fields(lines, text):

    fields = {}

    # -----------------------------------------------------
    # VENDOR
    # -----------------------------------------------------

    vendor = find_vendor(lines)

    if vendor:
        fields["Vendor"] = vendor

    # -----------------------------------------------------
    # INVOICE NUMBER
    # -----------------------------------------------------

    invoice = find_invoice_number(lines)

    if invoice:
        fields["Invoice Number"] = invoice

    # -----------------------------------------------------
    # INVOICE DATE
    # -----------------------------------------------------

    invoice_date = find_invoice_date(lines)

    if invoice_date:
        fields["Invoice Date"] = invoice_date

    # -----------------------------------------------------
    # PHONE
    # -----------------------------------------------------

    phone = find_phone(lines)

    if phone:
        fields["Phone"] = phone

    # -----------------------------------------------------
    # KNOWN FIELDS
    # -----------------------------------------------------

    known_fields = extract_known_fields(lines)

    for key, value in known_fields.items():

        if key not in fields:
            fields[key] = value

    # -----------------------------------------------------
    # DYNAMIC FIELDS
    # -----------------------------------------------------

    dynamic_fields = find_dynamic_labels(lines)

    for key, value in dynamic_fields.items():

        if key not in fields:
            fields[key] = value

    return fields


# =========================================================
# TABLE DETECTION
# =========================================================

HEADER_KEYWORDS = [

    "item",
    "description",
    "particulars",
    "details",
    "product",
    "service",
    "plan",
    "qty",
    "quantity",
    "unit",
    "rate",
    "price",
    "amount",
    "charge",
    "charges",
    "tax",
    "gst",
    "vat",
    "cgst",
    "sgst",
    "igst",
    "discount",
    "mrp",
    "fare",
    "distance",
    "usage",
    "hours",
    "days",
    "nights",
    "from",
    "to",
    "sku",
    "code",
    "room",
    "doctor",
    "medicine",
    "test",
    "bed"
]


# =========================================================
# SUMMARY KEYWORDS
# =========================================================

SUMMARY_KEYWORDS = [

    "subtotal",
    "sub total",
    "grand total",
    "amount payable",
    "amount due",
    "balance due",
    "net payable",
    "net amount",
    "total tax",
    "payment received",
    "amount paid",
    "final total",
    "thank you"
]


# =========================================================
# GROUP OCR/PDF WORDS INTO LINES
# =========================================================

def group_words_into_lines(words):

    words = sorted(
        words,
        key=lambda x: (
            x["top"],
            x["left"]
        )
    )

    lines = []

    for word in words:

        added = False

        for line in lines:

            if abs(
                word["top"]
                -
                line[0]["top"]
            ) <= 12:

                line.append(word)

                added = True

                break

        if not added:

            lines.append([word])

    for line in lines:

        line.sort(
            key=lambda x: x["left"]
        )

    return lines


# =========================================================
# CHECK HEADER
# =========================================================

def is_header_line(line_words):

    text = " ".join(
        word["text"].lower()
        for word in line_words
    )

    count = 0

    for keyword in HEADER_KEYWORDS:

        if re.search(
            r"\b"
            + re.escape(keyword)
            + r"\b",
            text
        ):

            count += 1

    return count >= 2


# =========================================================
# NORMALIZE COLUMN NAMES
# =========================================================

def normalize_column_name(name):

    mapping = {

        "s.no": "S.No.",
        "sno": "S.No.",
        "sr.no": "S.No.",
        "sr": "S.No.",

        "qty": "Quantity",
        "quantity": "Quantity",

        "rate": "Price",
        "price": "Price",
        "unit price": "Price",
        "unit rate": "Price",

        "description": "Item",
        "particulars": "Item",
        "details": "Item",
        "item": "Item",
        "product": "Item",
        "service": "Item",
        "plan": "Item",

        "amount": "Total",
        "total": "Total",

        "fare": "Fare",

        "tax": "Tax",
        "gst": "GST",

        "discount": "Discount",

        "from": "From",
        "to": "To",

        "distance": "Distance",

        "days": "Days",
        "nights": "Nights",
        "hours": "Hours"
    }

    clean_name = name.lower().strip()

    return mapping.get(
        clean_name,
        name.title()
    )


# =========================================================
# CREATE TABLE COLUMNS
# =========================================================

def create_columns(header_words):

    columns = []

    existing_names = set()

    for word in header_words:

        name = word["text"].strip()

        if not name:
            continue

        name = normalize_column_name(name)

        if name in existing_names:
            continue

        existing_names.add(name)

        columns.append({
            "name": name,
            "x": (
                word["left"]
                +
                word["width"] / 2
            )
        })

    columns.sort(
        key=lambda x: x["x"]
    )

    return columns


# =========================================================
# ASSIGN WORD TO COLUMN
# =========================================================

def assign_word_to_column(word, columns):

    word_x = (
        word["left"]
        +
        word["width"] / 2
    )

    columns = sorted(
        columns,
        key=lambda x: x["x"]
    )

    if len(columns) == 1:

        return columns[0]["name"]

    nearest_column = min(
        columns,
        key=lambda column:
        abs(
            word_x
            -
            column["x"]
        )
    )

    return nearest_column["name"]


# =========================================================
# DYNAMIC TABLE EXTRACTION
# =========================================================

def extract_dynamic_table(words):

    grouped_lines = group_words_into_lines(
        words
    )

    header_index = None

    # -----------------------------------------------------
    # FIND HEADER
    # -----------------------------------------------------

    for i, line in enumerate(grouped_lines):

        if is_header_line(line):

            header_index = i

            break

    if header_index is None:

        return [], []

    header_words = grouped_lines[
        header_index
    ]

    columns = create_columns(
        header_words
    )

    if not columns:

        return [], []

    rows = []

    # -----------------------------------------------------
    # READ ROWS
    # -----------------------------------------------------

    for line in grouped_lines[
        header_index + 1:
    ]:

        text = " ".join(
            word["text"]
            for word in line
        )

        lower_text = text.lower()

        # -------------------------------------------------
        # STOP WHEN SUMMARY STARTS
        # -------------------------------------------------

        if any(
            keyword in lower_text
            for keyword in SUMMARY_KEYWORDS
        ):
            break

        if not text.strip():
            continue

        row = {}

        for word in line:

            column = assign_word_to_column(
                word,
                columns
            )

            value = word["text"].strip()

            if not value:
                continue

            if column in row:

                row[column] += " " + value

            else:

                row[column] = value

        if row:

            rows.append(row)

    return columns, rows


# =========================================================
# CLEAN TABLE ROWS
# =========================================================

def clean_table_rows(rows):

    result = []

    for row in rows:

        new_row = {}

        for key, value in row.items():

            value = str(value).strip()

            if value:

                new_row[key] = value

        if new_row:

            result.append(new_row)

    return result


# =========================================================
# REMOVE EMPTY COLUMNS
# =========================================================

def remove_empty_columns(columns, rows):

    result = []

    for column in columns:

        name = column["name"]

        if any(
            row.get(name)
            for row in rows
        ):

            result.append(column)

    return result


# =========================================================
# OUTPUT FILE NAME
# =========================================================

def get_output_filename():

    if not os.path.exists(
        OUTPUT_FILE
    ):

        return OUTPUT_FILE

    name, extension = os.path.splitext(
        OUTPUT_FILE
    )

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    return (
        f"{name}_"
        f"{timestamp}"
        f"{extension}"
    )


# =========================================================
# AUTO COLUMN WIDTH
# =========================================================

def auto_width(sheet):

    for cells in sheet.columns:

        maximum = 0

        column = get_column_letter(
            cells[0].column
        )

        for cell in cells:

            if cell.value is not None:

                maximum = max(
                    maximum,
                    len(str(cell.value))
                )

        sheet.column_dimensions[
            column
        ].width = min(
            maximum + 2,
            50
        )


# =========================================================
# SAVE TO EXCEL
# =========================================================

def save_to_excel(
    fields,
    table_columns,
    table_rows,
    summary_values,
    output_file
):

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Bill Data"

    # =====================================================
    # TOP BILL INFORMATION
    # =====================================================

    current_row = 1

    if fields:

        field_headings = list(
            fields.keys()
        )

        # -------------------------------------------------
        # HEADINGS
        # -------------------------------------------------

        for column_number, heading in enumerate(
            field_headings,
            start=1
        ):

            cell = sheet.cell(
                row=current_row,
                column=column_number
            )

            cell.value = heading

            cell.font = Font(
                bold=True
            )

        # -------------------------------------------------
        # VALUES
        # -------------------------------------------------

        for column_number, heading in enumerate(
            field_headings,
            start=1
        ):

            sheet.cell(
                row=current_row + 1,
                column=column_number
            ).value = fields[heading]

        current_row += 3

    # =====================================================
    # ITEM / SERVICE TABLE
    # =====================================================

    table_headings = []

    if table_rows:

        # -------------------------------------------------
        # S.NO. FIRST
        # -------------------------------------------------

        table_headings.append(
            "S.No."
        )

        # -------------------------------------------------
        # OTHER COLUMNS
        # -------------------------------------------------

        for column in table_columns:

            name = column["name"]

            if name not in table_headings:

                table_headings.append(
                    name
                )

        # -------------------------------------------------
        # TABLE HEADER
        # -------------------------------------------------

        for column_number, heading in enumerate(
            table_headings,
            start=1
        ):

            cell = sheet.cell(
                row=current_row,
                column=column_number
            )

            cell.value = heading

            cell.font = Font(
                bold=True
            )

        # -------------------------------------------------
        # TABLE DATA
        # -------------------------------------------------

        for row_number, table_row in enumerate(
            table_rows,
            start=1
        ):

            excel_row = (
                current_row
                +
                row_number
            )

            for column_number, heading in enumerate(
                table_headings,
                start=1
            ):

                # S.No.
                if heading == "S.No.":

                    value = row_number

                else:

                    value = table_row.get(
                        heading,
                        ""
                    )

                sheet.cell(
                    row=excel_row,
                    column=column_number
                ).value = value

        # -------------------------------------------------
        # MOVE BELOW TABLE
        # -------------------------------------------------

        current_row += (
            len(table_rows)
            +
            2
        )

    # =====================================================
    # BILL SUMMARY
    # =====================================================

    if summary_values:

        # -------------------------------------------------
        # BILL SUMMARY
        # -------------------------------------------------

        sheet.cell(
            row=current_row,
            column=1
        ).value = "BILL SUMMARY"

        sheet.cell(
            row=current_row,
            column=1
        ).font = Font(
            bold=True
        )

        # -------------------------------------------------
        # FIND TOTAL COLUMN
        # -------------------------------------------------

        total_column = None

        if table_rows:

            for column_number, heading in enumerate(
                table_headings,
                start=1
            ):

                if heading.lower() in [
                    "total",
                    "amount",
                    "fare"
                ]:

                    total_column = column_number

                    break

        # -------------------------------------------------
        # IF TOTAL COLUMN NOT FOUND
        # -------------------------------------------------

        if total_column is None:

            if table_headings:

                total_column = len(
                    table_headings
                )

            else:

                total_column = 2

        # -------------------------------------------------
        # SUMMARY START
        # -------------------------------------------------

        current_row += 1

        # -------------------------------------------------
        # SUMMARY DATA
        # -------------------------------------------------

        for key in [

            "Subtotal",
            "Discount",
            "Tax",
            "Other Charges",
            "Final Total"

        ]:

            value = summary_values.get(
                key
            )

            if not value:
                continue

            # -------------------------------------------------
            # LABEL
            # -------------------------------------------------

            sheet.cell(
                row=current_row,
                column=1
            ).value = key

            sheet.cell(
                row=current_row,
                column=1
            ).font = Font(
                bold=True
            )

            # -------------------------------------------------
            # VALUE UNDER TOTAL
            # -------------------------------------------------

            sheet.cell(
                row=current_row,
                column=total_column
            ).value = value

            sheet.cell(
                row=current_row,
                column=total_column
            ).font = Font(
                bold=True
            )

            current_row += 1

    # =====================================================
    # EXCEL FORMATTING
    # =====================================================

    auto_width(sheet)

    sheet.freeze_panes = "A2"

    workbook.save(
        output_file
    )


# =========================================================
# MAIN
# =========================================================

def main():

    print(
        "\n========== BILL OCR ==========\n"
    )

    # -----------------------------------------------------
    # CHECK INPUT FILE
    # -----------------------------------------------------

    if not os.path.exists(INPUT_FILE):

        print(
            f"ERROR: File not found: {INPUT_FILE}"
        )

        return

    # -----------------------------------------------------
    # READ BILL
    # -----------------------------------------------------

    text, words = extract_bill_data(
        INPUT_FILE
    )

    # -----------------------------------------------------
    # CLEAN TEXT
    # -----------------------------------------------------

    lines = clean_text(text)

    # -----------------------------------------------------
    # PRINT OCR / PDF TEXT
    # -----------------------------------------------------

    print(
        "\n========== OCR / PDF TEXT ==========\n"
    )

    print(text)

    # -----------------------------------------------------
    # EXTRACT BILL INFORMATION
    # -----------------------------------------------------

    fields = create_bill_fields(
        lines,
        text
    )

    # -----------------------------------------------------
    # EXTRACT ITEM / SERVICE TABLE
    # -----------------------------------------------------

    table_columns, table_rows = (
        extract_dynamic_table(words)
    )

    # -----------------------------------------------------
    # CLEAN TABLE
    # -----------------------------------------------------

    table_rows = clean_table_rows(
        table_rows
    )

    table_columns = remove_empty_columns(
        table_columns,
        table_rows
    )

    # -----------------------------------------------------
    # BILL SUMMARY
    # -----------------------------------------------------

    summary_values = {

        "Subtotal":
            find_subtotal(lines),

        "Discount":
            find_discount(lines),

        "Tax":
            find_tax(lines),

        "Other Charges":
            find_other_charges(lines),

        "Final Total":
            find_total(lines)
    }

    # -----------------------------------------------------
    # REMOVE EMPTY SUMMARY FIELDS
    # -----------------------------------------------------

    summary_values = {

        key: value

        for key, value
        in summary_values.items()

        if value
    }

    # -----------------------------------------------------
    # OUTPUT FILE
    # -----------------------------------------------------

    output_file = get_output_filename()

    # -----------------------------------------------------
    # SAVE EXCEL
    # -----------------------------------------------------

    save_to_excel(

        fields,

        table_columns,

        table_rows,

        summary_values,

        output_file
    )

    # =====================================================
    # TERMINAL OUTPUT
    # =====================================================

    print(
        "\n========== DETECTED DATA ==========\n"
    )

    for key, value in fields.items():

        print(
            f"{key}: {value}"
        )

    print(
        "\n========== ITEMS / SERVICES ==========\n"
    )

    for number, row in enumerate(
        table_rows,
        start=1
    ):

        print(
            number,
            row
        )

    print(
        "\n========== BILL SUMMARY ==========\n"
    )

    for key, value in summary_values.items():

        print(
            f"{key}: {value}"
        )

    print(
        "\n===================================="
    )

    print(
        "Excel created:"
    )

    print(
        output_file
    )

    print(
        "====================================\n"
    )


# =========================================================
# START PROGRAM
# =========================================================

if __name__ == "__main__":
    main()
